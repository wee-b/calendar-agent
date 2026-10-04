import json
import unittest
from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
import httpx
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from fastapi import HTTPException

from app.all_graph.nodes.plan_node import PlanNode
from app.schemas.plan import today
from app.schemas.statemachine.flow import PendingTask, UserSignal
from app.service.documents import (DocumentUploadService, DocumentParseService,
                                   DocumentSearchService, DocumentManagementService, extract_chunks)
from app.repository.qdrant import QdrantRepository
from app.schemas.rag import QdrantPayload, QdrantPoint


def valid_plan():
    day = today()
    return json.dumps({"goal": "复习", "startDate": str(day), "endDate": str(day),
                       "todos": [{"title": "复习", "startDate": str(day),
                                  "endDate": str(day), "weekDays": [day.isoweekday()]}]})


class DocumentTests(unittest.IsolatedAsyncioTestCase):
    def test_supported_file_formats_extract_real_text(self):
        phrase = "每天复习十个单词并练习听力。"
        for extension, encoding in (("txt", "utf-8"), ("txt", "gb18030"), ("md", "utf-8")):
            raw = (f"# 学习目标\n{phrase}" if extension == "md" else phrase).encode(encoding)
            chunks = extract_chunks(f"sample.{extension}", raw)
            self.assertIn(phrase, chunks[0].text)
            if extension == "md":
                self.assertEqual("学习目标", chunks[0].section)

        document = Document()
        document.add_paragraph(phrase)
        document.add_table(rows=1, cols=1).cell(0, 0).text = "每周进行一次模拟练习。"
        docx = BytesIO()
        document.save(docx)
        chunks = extract_chunks("sample.docx", docx.getvalue())
        self.assertIn(phrase, chunks[0].text)
        self.assertIn("每周进行一次模拟练习", chunks[0].text)

        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 72 720 Td (Daily review practice and vocabulary study.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        pdf = BytesIO()
        writer.write(pdf)
        chunks = extract_chunks("sample.pdf", pdf.getvalue())
        self.assertIn("Daily review practice", chunks[0].text)

        blank_pdf = BytesIO()
        blank = PdfWriter()
        blank.add_blank_page(width=612, height=792)
        blank.write(blank_pdf)
        with self.assertRaisesRegex(HTTPException, "没有可索引的文本"):
            extract_chunks("scanned.pdf", blank_pdf.getvalue())

    def test_extract_deduplicates_identical_paragraphs(self):
        text = "每天复习十个单词并进行听力练习。\n\n每天复习十个单词并进行听力练习。"
        chunks = extract_chunks("notes.md", text.encode())
        self.assertEqual(1, len(chunks))

    async def test_search_rejects_other_users_file_before_vector_query(self):
        repo = SimpleNamespace(ready_files=AsyncMock(return_value=[]))
        embedding = SimpleNamespace(embed=AsyncMock(return_value=[1.0]))
        qdrant = SimpleNamespace(search_documents=AsyncMock())
        service = DocumentSearchService(repository=repo, qdrant=qdrant, embedding=embedding)
        with self.assertRaises(HTTPException):
            await service.search(7, [42], "复习")
        embedding.embed.assert_not_awaited()
        qdrant.search_documents.assert_not_awaited()

    async def test_document_search_logs_query_and_returned_hits(self):
        repo = SimpleNamespace(ready_files=AsyncMock(return_value=[SimpleNamespace(file_id=42)]))
        embedding = SimpleNamespace(embed=AsyncMock(return_value=[1.0]))
        hit = QdrantPoint(id="point-1", score=0.87, payload=QdrantPayload(
            user_id=7, file_id=42, source="notes.md", section="复习", text="每天背单词"))
        qdrant = SimpleNamespace(search_documents=AsyncMock(return_value=[hit]))
        service = DocumentSearchService(repository=repo, qdrant=qdrant, embedding=embedding)
        with self.assertLogs("app.service.documents", level="INFO") as captured:
            result = await service.search(7, [42, 42], "明天复习")
        self.assertEqual([hit], result)
        self.assertIn("文件ID=[42]", captured.output[0])
        self.assertIn("查询='明天复习'", captured.output[0])
        self.assertIn("命中数=1", captured.output[1])
        self.assertIn("点ID=point-1", captured.output[2])
        self.assertIn("内容='每天背单词'", captured.output[2])

    async def test_qdrant_query_filters_user_and_selected_files(self):
        captured = {}

        def handler(request):
            captured.update(json.loads(request.content))
            return httpx.Response(200, json={"result": {"points": []}})

        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            await QdrantRepository(client, "user_documents").search_documents([1.0], 7, [42], 3)
        self.assertEqual(7, captured["filter"]["must"][0]["match"]["value"])
        self.assertEqual([42], captured["filter"]["must"][1]["match"]["any"])

    async def test_upload_only_stores_file_then_parse_indexes_chunks(self):
        class FakeMinio:
            def __init__(self):
                self.created = []
                self.objects = []

            def bucket_exists(self, name):
                return bool(self.created)

            def make_bucket(self, name):
                self.created.append(name)

            def put_object(self, bucket, key, stream, size, content_type):
                self.objects.append((bucket, key, stream.read(), size, content_type))

            def get_object(self, bucket, key):
                payload = next(item[2] for item in self.objects if item[0] == bucket and item[1] == key)
                stream = BytesIO(payload)
                stream.release_conn = lambda: None
                return stream

        data = "每天背单词并练习听力。".encode()
        entry = SimpleNamespace(file_id=5, file_name="notes.txt", size_bytes=len(data),
                                sha256=sha256(data).hexdigest(),
                                status="uploading", error_message=None, create_time=None,
                                object_key="", object_bucket="calendar-documents")

        async def create_or_get(**kwargs):
            entry.object_key = kwargs["object_key"]
            return entry, True

        repo = SimpleNamespace(create_or_get=AsyncMock(side_effect=create_or_get),
                               finish_upload=AsyncMock(), fail_upload=AsyncMock(),
                               claim_parse=AsyncMock(return_value=(entry, True, "uploaded")),
                               clear_chunks=AsyncMock(), finish_parse=AsyncMock(), fail=AsyncMock())
        qdrant = SimpleNamespace(collection_info=AsyncMock(return_value=None),
                                 create_collection=AsyncMock(), upsert=AsyncMock(),
                                 delete_file=AsyncMock())
        minio = FakeMinio()
        embedding = SimpleNamespace(embed=AsyncMock(return_value=[1.0, 0.0]))
        service = DocumentUploadService(repository=repo, qdrant=qdrant,
                                  embedding=embedding,
                                  minio_client=minio)
        result = await service.upload(7, "notes.txt", data, "text/plain")
        self.assertEqual("uploaded", result["status"])
        self.assertEqual([service.settings.minio_bucket], minio.created)
        self.assertEqual(data, minio.objects[0][2])
        repo.finish_upload.assert_awaited_once_with(5)
        qdrant.upsert.assert_not_awaited()
        embedding.embed.assert_not_awaited()

        parser = DocumentParseService(repository=repo, qdrant=qdrant,
                                     embedding=embedding, minio_client=minio)
        result = await parser.parse(7, 5)
        self.assertEqual("ready", result["status"])
        points = qdrant.upsert.call_args.args[0]
        self.assertEqual(7, points[0].payload.user_id)
        self.assertEqual(5, points[0].payload.file_id)
        repo.finish_parse.assert_awaited_once()

        repo.claim_parse.return_value = (entry, True, "ready")
        result = await parser.parse(7, 5, force=True)
        self.assertEqual("ready", result["status"])
        repo.claim_parse.assert_awaited_with(7, 5, force=True)
        self.assertEqual(2, repo.clear_chunks.await_count)
        self.assertEqual(2, qdrant.delete_file.await_count)

    async def test_parse_rejects_unowned_file_before_storage_or_embedding(self):
        repo = SimpleNamespace(claim_parse=AsyncMock(return_value=(None, False, None)))
        qdrant = SimpleNamespace(upsert=AsyncMock())
        embedding = SimpleNamespace(embed=AsyncMock())
        service = DocumentParseService(repository=repo, qdrant=qdrant, embedding=embedding)
        with self.assertRaises(HTTPException) as raised:
            await service.parse(7, 42)
        self.assertEqual(404, raised.exception.status_code)
        embedding.embed.assert_not_awaited()
        qdrant.upsert.assert_not_awaited()

    async def test_reparse_embedding_failure_keeps_old_index_ready(self):
        data = "每天复习十个单词并练习听力。".encode()
        entry = SimpleNamespace(file_id=5, file_name="notes.txt", size_bytes=len(data),
                                sha256=sha256(data).hexdigest(), object_key="source",
                                object_bucket="calendar-documents",
                                status="processing", error_message=None)
        stream = BytesIO(data)
        stream.release_conn = lambda: None
        minio = SimpleNamespace(get_object=lambda *_: stream)
        repo = SimpleNamespace(claim_parse=AsyncMock(return_value=(entry, True, "ready")),
                               fail=AsyncMock(), clear_chunks=AsyncMock())
        qdrant = SimpleNamespace(delete_file=AsyncMock())
        embedding = SimpleNamespace(embed=AsyncMock(side_effect=RuntimeError("embedding unavailable")))
        service = DocumentParseService(repository=repo, qdrant=qdrant,
                                       embedding=embedding, minio_client=minio)
        with self.assertRaisesRegex(RuntimeError, "embedding unavailable"):
            await service.parse(7, 5, force=True)
        qdrant.delete_file.assert_not_awaited()
        repo.clear_chunks.assert_not_awaited()
        repo.fail.assert_awaited_once_with(5, "embedding unavailable", restore_ready=True)

    async def test_delete_parse_preserves_source_and_delete_source_removes_all(self):
        entry = SimpleNamespace(file_id=5, user_id=7, object_key="users/7/source.txt",
                                object_bucket="calendar-documents",
                                file_name="source.txt", size_bytes=20, create_time=None,
                                status="ready", error_message=None)
        repo = SimpleNamespace(claim_delete_parse=AsyncMock(return_value=(entry, True)),
                               complete_delete_parse=AsyncMock(), fail_delete=AsyncMock(),
                               claim_delete_source=AsyncMock(return_value=(entry, True)),
                               complete_delete_source=AsyncMock())
        qdrant = SimpleNamespace(collection_info=AsyncMock(return_value=object()),
                                 delete_file=AsyncMock())
        minio = SimpleNamespace(remove_object=unittest.mock.Mock())
        service = DocumentManagementService(repository=repo, qdrant=qdrant,
                                            minio_client=minio)
        result = await service.delete_parse(7, 5)
        self.assertEqual("uploaded", result["status"])
        repo.complete_delete_parse.assert_awaited_once_with(5)
        minio.remove_object.assert_not_called()

        await service.delete_source(7, 5)
        minio.remove_object.assert_called_once_with(service.settings.minio_bucket, entry.object_key)
        repo.complete_delete_source.assert_awaited_once_with(5)
        self.assertEqual(2, qdrant.delete_file.await_count)

    async def test_delete_source_rejects_other_users_file(self):
        repo = SimpleNamespace(claim_delete_source=AsyncMock(return_value=(None, False)))
        qdrant = SimpleNamespace(delete_file=AsyncMock())
        minio = SimpleNamespace(remove_object=unittest.mock.Mock())
        service = DocumentManagementService(repository=repo, qdrant=qdrant,
                                            minio_client=minio)
        with self.assertRaises(HTTPException) as raised:
            await service.delete_source(7, 42)
        self.assertEqual(404, raised.exception.status_code)
        qdrant.delete_file.assert_not_awaited()
        minio.remove_object.assert_not_called()

    async def test_plan_searches_only_when_document_referenced(self):
        hit = SimpleNamespace(payload=SimpleNamespace(source="notes.md", section="正文", text="每天背单词"))
        documents = SimpleNamespace(search=AsyncMock(return_value=[hit]))
        planning = SimpleNamespace(memories=AsyncMock(return_value=[]),
                                   save_draft=AsyncMock(return_value=SimpleNamespace(draft_id=3)))
        model = SimpleNamespace(complete=AsyncMock(return_value=valid_plan()))
        context = SimpleNamespace(use_model=lambda *_: None,
                                  add_step=AsyncMock(return_value=1), complete_step=AsyncMock(),
                                  set_agent_narration=AsyncMock())
        node = PlanNode(planning, model, documents)
        state = {"user_id": 7, "session_id": "s", "message": "明天复习",
                 "pending": PendingTask(), "signal": UserSignal.NEW_PLAN}
        await node(state, context)
        documents.search.assert_not_awaited()
        self.assertNotIn("每天背单词", model.complete.call_args.args[0][1]["content"])
        state["document_ids"] = [42]
        await node(state, context)
        documents.search.assert_awaited_once_with(7, [42], "明天复习")
        context.add_step.assert_any_await("tool", "检索引用文档")
        context.complete_step.assert_awaited_once_with(1)
        self.assertIn("每天背单词", model.complete.call_args.args[0][1]["content"])
